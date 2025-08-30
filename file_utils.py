from pathlib import Path
from datetime import datetime
import hashlib
from PIL import Image
from PIL.ExifTags import TAGS
import pyexiv2  # pip install pyexiv2
import pillow_heif  # pip install pillow-heif
import fitz  # pip install pymupdf
import pdfplumber  # pip install pdfplumber
from docx import Document  # pip install python-docx
import openpyxl  # pip install openpyxl
from pptx import Presentation  # pip install python-pptx
import mutagen  # pip install mutagen
import cv2  # pip install opencv-python
import zipfile
import tarfile
import py7zr  # pip install py7zr
import os
import magic
from pathlib import Path
import ffmpeg  # pip install ffmpeg-python


def read_file_attributes_pathlib(filepath):
    """Read file attributes using pathlib (Python 3.4+)"""
    try:
        path = Path(filepath)

        if not path.exists():
            print(f"File {filepath} not found")
            return

        file_stats = path.stat()

        print(f"File: {filepath}")
        print("-" * 40)
        print(f"Name: {path.name}")
        print(f"Suffix: {path.suffix}")
        print(f"Parent directory: {path.parent}")
        print(f"Absolute path: {path.absolute()}")
        print(f"Size: {file_stats.st_size} bytes")
        print(f"Modified: {datetime.fromtimestamp(file_stats.st_mtime)}")
        print(f"Is file: {path.is_file()}")
        print(f"Is directory: {path.is_dir()}")
        print(f"Is symlink: {path.is_symlink()}")

        # Human-readable size
        size_mb = file_stats.st_size / (1024 * 1024)
        if size_mb > 1:
            print(f"Size (MB): {size_mb:.2f} MB")
        else:
            size_kb = file_stats.st_size / 1024
            print(f"Size (KB): {size_kb:.2f} KB")
    except FileNotFoundError as e:
        print(f"File not found: {e}")
    except PermissionError as e:
        print(f"Permission denied: {e}")
    except OSError as e:
        print(f"OS error: {e}")
    except Exception as e:
        print(f"Error reading file attributes: {e}")


def _extract_image_metadata(filepath):
    # Basic image info with Pillow
    with Image.open(filepath) as img:
        basic = {"size": img.size, "mode": img.mode, "format": img.format}

    # Detailed metadata with pyexiv2
    img_meta = pyexiv2.Image(filepath)
    exif = img_meta.read_exif()
    iptc = img_meta.read_iptc()
    xmp = img_meta.read_xmp()

    return {"basic": basic, "exif": exif, "iptc": iptc, "xmp": xmp}


def _extract_pdf_metadata(filepath):
    # Using pymupdf for comprehensive metadata
    doc = fitz.open(filepath)
    metadata = doc.metadata

    # Additional info
    info = {
        "page_count": doc.page_count,
        "is_pdf": doc.is_pdf,
        "is_encrypted": doc.needs_pass,
        "metadata": metadata,
    }

    # Text extraction with pdfplumber
    with pdfplumber.open(filepath) as pdf:
        text = ""
        for page in pdf.pages:
            text += page.extract_text() or ""

    info["extracted_text"] = text
    doc.close()
    return info


def _extract_docx_metadata(filepath):
    doc = Document(filepath)
    props = doc.core_properties

    return {
        "title": props.title,
        "author": props.author,
        "subject": props.subject,
        "created": props.created,
        "modified": props.modified,
        "last_modified_by": props.last_modified_by,
        "revision": props.revision,
    }


def _extract_xlsx_metadata(filepath):
    wb = openpyxl.load_workbook(filepath)
    props = wb.properties

    return {
        "title": props.title,
        "creator": props.creator,
        "created": props.created,
        "modified": props.modified,
        "sheet_names": wb.sheetnames,
    }


def _extract_pptx_metadata(filepath):
    prs = Presentation(filepath)
    props = prs.core_properties

    return {
        "title": props.title,
        "author": props.author,
        "created": props.created,
        "modified": props.modified,
        "slide_count": len(prs.slides),
    }


def _extract_audio_metadata(filepath):
    audio_file = mutagen.File(filepath)

    if audio_file is None:
        return None

    metadata = dict(audio_file)

    # Common fields (vary by format)
    common = {}
    if hasattr(audio_file, "info"):
        common.update(
            {
                "length": audio_file.info.length,
                "bitrate": getattr(audio_file.info, "bitrate", None),
                "sample_rate": getattr(audio_file.info, "sample_rate", None),
                "channels": getattr(audio_file.info, "channels", None),
            }
        )

    return {"raw_metadata": metadata, "common": common}


def _extract_video_metadata(filepath):
    # Using OpenCV
    cap = cv2.VideoCapture(filepath)

    metadata = {
        "frame_count": int(cap.get(cv2.CAP_PROP_FRAME_COUNT)),
        "fps": cap.get(cv2.CAP_PROP_FPS),
        "width": int(cap.get(cv2.CAP_PROP_FRAME_WIDTH)),
        "height": int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT)),
        "duration": cap.get(cv2.CAP_PROP_FRAME_COUNT) / cap.get(cv2.CAP_PROP_FPS),
    }

    cap.release()
    return metadata


# Alternative with ffmpeg
# def _extract_video_metadata_ffmpeg(filepath):
#     probe = ffmpeg.probe(filepath)
#     return {"format": probe["format"], "streams": probe["streams"]}


# def _extract_archive_metadata(filepath):
#     ext = filepath.lower().split(".")[-1]

#     if ext == "zip":
#         with zipfile.ZipFile(filepath, "r") as zf:
#             return {
#                 "file_count": len(zf.namelist()),
#                 "files": zf.namelist(),
#                 "compressed_size": sum(info.compress_size for info in zf.infolist()),
#                 "uncompressed_size": sum(info.file_size for info in zf.infolist()),
#             }

#     elif ext in ["tar", "gz", "bz2", "xz"]:
#         with tarfile.open(filepath, "r") as tf:
#             return {"file_count": len(tf.getnames()), "files": tf.getnames()}

#     elif ext == "7z":
#         with py7zr.SevenZipFile(filepath, mode="r") as z7:
#             return {"file_count": len(z7.getnames()), "files": z7.getnames()}


def extract_all_metadata(filepath):
    filepath = Path(filepath)
    mime_type = magic.from_file(str(filepath), mime=True)

    # Basic file stats
    stat = filepath.stat()
    result = {
        "basic": {
            "filename": filepath.name,
            "size": stat.st_size,
            "mime_type": mime_type,
            "created": stat.st_ctime,
            "modified": stat.st_mtime,
        }
    }

    # Route to appropriate extractor
    if mime_type.startswith("image/"):
        result["image"] = _extract_image_metadata(filepath)
    elif mime_type == "application/pdf":
        result["pdf"] = _extract_pdf_metadata(filepath)
    elif "document" in mime_type or filepath.suffix in [".docx", ".xlsx", ".pptx"]:
        if filepath.suffix == ".docx":
            result["document"] = _extract_docx_metadata(filepath)
        elif filepath.suffix == ".xlsx":
            result["document"] = _extract_xlsx_metadata(filepath)
        elif filepath.suffix == ".pptx":
            result["document"] = _extract_pptx_metadata(filepath)
    elif mime_type.startswith("audio/"):
        result["audio"] = _extract_audio_metadata(filepath)
    elif mime_type.startswith("video/"):
        result["video"] = _extract_video_metadata(filepath)
    # elif mime_type in ["application/zip", "application/x-tar"]:
    #     result["archive"] = _extract_archive_metadata(filepath)

    return result
